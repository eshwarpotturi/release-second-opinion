import openpyxl, glob, re, csv, warnings
warnings.filterwarnings("ignore")
rows=[]; logs=[]
for f in sorted(glob.glob("raw_new/RIA-*.xlsx")):
    rel=re.search(r"RIA-(\w+)\.xlsx",f).group(1)
    wb=openpyxl.load_workbook(f,read_only=True)
    ws=[w for w in wb.worksheets if w.title in("RIA","Vault RIA","Veeva Vault RIA")][0]
    it=ws.iter_rows(values_only=True); h=[str(c).strip() if c else "" for c in next(it)]
    ix={k:h.index(k) for k in ["Application Family","Application","Name","Enablement Setting","GxP Risk","Default Impact","Feature Description"]}
    fam=None; n=0
    for r in it:
        if not r[ix["Name"]]: continue
        risk=str(r[ix["GxP Risk"]] or "").strip()
        if risk not in("High","Medium","Low","N/A"): continue
        fam=r[ix["Application Family"]] or fam
        rows.append([rel,fam,r[ix["Application"]],str(r[ix["Name"]]).replace("\n"," ").strip(),risk,r[ix["Enablement Setting"]],r[ix["Default Impact"]],str(r[ix["Feature Description"]] or "").replace("\n"," ").strip(),""]); n+=1
    cl=wb["Change Log"]
    txt="\n".join(str(c) for row in cl.iter_rows(values_only=True) for c in row if c)
    open(f"data/changelog_{rel}.txt","w").write(txt)
    print(rel,n)
with open("data/features_new.csv","w",newline="") as fh:
    w=csv.writer(fh); w.writerow("release section application feature risk enablement default_impact description page".split()); w.writerows(rows)
